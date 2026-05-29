import argparse
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from phase_evaluator import DEFAULT_RESULT_DIR, ROOT


GUIDE_ROWS = [
    ["항목", "설명"],
    ["파일 목적", "MP alloy 5000개 결과에서 MP 준안정상 구간을 포함해 모델 성능과 예측 분포를 요약한 benchmark workbook입니다."],
    ["엄격 binary 지표", "MP stable_single과 MP unstable_multi_risk만 사용해 single/multi 성능을 계산합니다. low/high metastable은 정답이 애매하므로 binary 성능 계산에서 제외합니다."],
    ["MP 안정성 구간 개수", "5000개 전체가 MP 기준 안정성 구간별로 몇 개인지 보여줍니다."],
    ["MP구간별 예측분포", "MP 안정성 구간마다 모델이 싱글/멀티/다른 결정구조 가능으로 어떻게 예측했는지 보여줍니다."],
    ["", ""],
    ["MP 안정성 구간", "의미"],
    ["stable_single", "MP energy above hull <= 0.001 eV/atom. MP 기준 ground-state stable입니다."],
    ["low_metastable", "0.001 < MP energy above hull <= 0.03 eV/atom. 낮은 준안정 영역으로 single phase 가능성을 배제하기 어렵습니다."],
    ["high_metastable", "0.03 < MP energy above hull <= 0.10 eV/atom. 조건 의존성이 크고 multi phase 위험도 커지는 영역입니다."],
    ["unstable_multi_risk", "MP energy above hull > 0.10 eV/atom. MP 기준 다상 분해 위험이 큰 영역입니다."],
    ["", ""],
    ["시트명", "의미"],
    ["엄격 binary 지표", "stable_single vs unstable_multi_risk만 비교한 precision/recall/F1/accuracy 지표입니다."],
    ["MP 안정성 구간 개수", "stable_single, low_metastable, high_metastable, unstable_multi_risk의 개수와 비율입니다."],
    ["MP구간별 예측분포", "각 MP 안정성 구간 안에서 모델 예측결과가 어떻게 분포하는지 보여줍니다."],
    ["", ""],
    ["주요 지표", "의미"],
    ["accuracy", "엄격 binary subset에서 전체 정답률입니다."],
    ["balanced_accuracy", "single recall과 multi recall의 평균입니다. 클래스 불균형에 더 안전합니다."],
    ["single_precision", "모델이 single이라고 한 것 중 MP stable_single인 비율입니다."],
    ["single_recall", "MP stable_single 중 모델이 single로 잡은 비율입니다."],
    ["multi_precision", "모델이 multi라고 한 것 중 MP unstable_multi_risk인 비율입니다."],
    ["multi_recall", "MP unstable_multi_risk 중 모델이 multi로 잡은 비율입니다."],
    ["tp_single", "MP stable_single이고 모델도 single로 예측한 수입니다."],
    ["fn_single", "MP stable_single인데 모델이 multi로 예측한 수입니다."],
    ["fp_single", "MP unstable_multi_risk인데 모델이 single로 예측한 수입니다."],
    ["tn_single", "MP unstable_multi_risk이고 모델도 multi로 예측한 수입니다."],
]


def style_guide_sheet(ws):
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 120
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    section_fill = PatternFill("solid", fgColor="E2F0D9")

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row_idx in [7, 13, 18]:
        for cell in ws[row_idx]:
            cell.font = Font(bold=True)
            cell.fill = section_fill


def parse_args():
    parser = argparse.ArgumentParser(description="Add guide sheet to benchmark workbook.")
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_RESULT_DIR / "benchmark_with_metastable_20260513_150336.xlsx",
    )
    parser.add_argument("--result-dir", type=Path, default=DEFAULT_RESULT_DIR)
    return parser.parse_args()


def main():
    args = parse_args()
    input_path = args.input if args.input.is_absolute() else ROOT / args.input
    result_dir = args.result_dir if args.result_dir.is_absolute() else ROOT / args.result_dir

    wb = load_workbook(input_path)
    if "안내" in wb.sheetnames:
        del wb["안내"]
    ws = wb.create_sheet("안내", 0)
    for row in GUIDE_ROWS:
        ws.append(row)
    style_guide_sheet(ws)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = result_dir / f"{input_path.stem}_guide_{timestamp}.xlsx"
    wb.save(output_path)
    print(f"Benchmark workbook with guide written to {output_path}")


if __name__ == "__main__":
    main()
