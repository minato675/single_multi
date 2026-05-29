import argparse
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from phase_evaluator import DEFAULT_RESULT_DIR, ROOT


GUIDE_ROWS = [
    ["항목", "설명"],
    ["파일 목적", "MP alloy 5000개 구조에 대해 CHGNet-only hull 기반 single/multi 예측과 MP 실제 평가를 함께 정리한 workbook입니다."],
    ["예측 기준", "CHGNet으로 MP competing phase와 후보 phase의 energy를 모두 다시 계산해 CHGNet-only phase diagram을 만들고, 후보의 energy above hull을 기준으로 분류합니다."],
    ["기본 single threshold", "CHGNet-only energy above hull <= 0.03 eV/atom이면 '싱글 후보'로 분류합니다."],
    ["다상 분해 가능", "예상 decomposition product가 후보 조성과 다른 여러 phase 조합으로 나타나는 경우입니다."],
    ["동일 조성 구조경쟁", "예상 phase가 후보와 같은 조성 하나로 나타나는 경우입니다. multi phase 분해라기보다 다른 polymorph/결정구조가 더 안정하다는 뜻입니다."],
    ["MP stable_single", "MP energy above hull <= 0.001 eV/atom. MP 기준 ground-state stable에 해당합니다."],
    ["MP low_metastable", "0.001 < MP energy above hull <= 0.03 eV/atom. 낮은 준안정 영역입니다."],
    ["MP high_metastable", "0.03 < MP energy above hull <= 0.10 eV/atom. 조건 의존성이 큰 준안정 영역입니다."],
    ["MP unstable_multi_risk", "MP energy above hull > 0.10 eV/atom. MP 기준 다상 분해 위험이 큰 영역입니다."],
    ["주의", "이 결과는 CHGNet 기반 1차 screening입니다. 최종 후보는 DFT relaxation과 실험 조건 검토가 필요합니다."],
    ["", ""],
    ["시트명", "의미"],
    ["MP 평가 전체", "5000개 전체 결과입니다. 모델 예측과 MP 평가 컬럼이 모두 들어 있습니다."],
    ["멀티 가능성 높음", "모델이 다상 분해 가능성이 높다고 예측한 항목입니다."],
    ["싱글 후보", "모델이 현재 CIF 구조를 single phase 후보로 예측한 항목입니다."],
    ["싱글(다른 결정구조 가능)", "후보 조성 자체는 single로 가능하지만, 현재 CIF와 다른 결정구조/polymorph가 더 안정할 수 있다고 예측한 항목입니다."],
    ["MP_high_metastable", "MP 기준 high metastable 영역에 있는 항목입니다."],
    ["MP_low_metastable", "MP 기준 low metastable 영역에 있는 항목입니다."],
    ["MP_stable_single", "MP 기준 stable single로 볼 수 있는 항목입니다."],
    ["MP_unstable_multi_risk", "MP 기준 multi phase 분해 위험이 큰 항목입니다."],
    ["", ""],
    ["주요 컬럼", "의미"],
    ["cif파일 제목", "MP material id 기반의 CIF 파일명입니다."],
    ["화학식", "후보 구조의 reduced formula입니다."],
    ["결정구조", "후보 구조의 space group symbol입니다."],
    ["단상 점수", "CHGNet-only E_hull 기반 single score입니다. 1에 가까울수록 single 후보입니다."],
    ["다상 점수", "1 - 단상 점수입니다."],
    ["예측결과(멀티인지 싱글인지)", "모델의 최종 분류 결과입니다."],
    ["경쟁 유형", "다상 분해 가능 또는 동일 조성 구조경쟁을 표시합니다."],
    ["예상되는 멀티 phase 조성들", "decomposition product 조합과 비율입니다."],
    ["예상 phase 결정구조들", "각 decomposition product의 space group과 MP entry id입니다."],
    ["MP_energy_above_hull_eV_per_atom", "MP의 DFT 기반 energy above hull입니다."],
    ["MP_stability_class", "MP E_hull 기준 안정성 구간입니다."],
]


def style_guide_sheet(ws):
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 120
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    section_fill = PatternFill("solid", fgColor="E2F0D9")

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row_idx in [13, 24]:
        for cell in ws[row_idx]:
            cell.font = Font(bold=True)
            cell.fill = section_fill


def parse_args():
    parser = argparse.ArgumentParser(description="Add guide sheet to an existing workbook.")
    parser.add_argument(
        "--input",
        type=Path,
        default=DEFAULT_RESULT_DIR / "phase_result_mp_alloys_with_MP_eval_20260513_145518.xlsx",
    )
    parser.add_argument("--result-dir", type=Path, default=DEFAULT_RESULT_DIR)
    return parser.parse_args()


def main():
    args = parse_args()
    input_path = args.input if args.input.is_absolute() else ROOT / args.input
    result_dir = args.result_dir if args.result_dir.is_absolute() else ROOT / args.result_dir

    wb = load_workbook(input_path)
    if "안내" in wb.sheetnames:
        del wb["안내"]
    ws = wb.create_sheet("안내", 0)
    for row in GUIDE_ROWS:
        ws.append(row)
    style_guide_sheet(ws)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = result_dir / f"{input_path.stem}_guide_{timestamp}.xlsx"
    wb.save(output_path)
    print(f"Workbook with guide written to {output_path}")


if __name__ == "__main__":
    main()
