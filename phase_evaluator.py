import argparse
import csv
import math
import os
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pymatgen.analysis.phase_diagram import PhaseDiagram
from pymatgen.core import Structure
from pymatgen.entries.computed_entries import ComputedStructureEntry
from pymatgen.symmetry.analyzer import SpacegroupAnalyzer


ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT_DIR = ROOT / "input_cif.dir"
DEFAULT_RESULT_DIR = ROOT / "result.dir"

KB_EV_PER_K = 8.617333262145e-5
DEFAULT_ENTROPY_TEMPERATURE_K = 1000.0
DEFAULT_ENTROPY_WEIGHT = 1.0

A_SITE_ELEMENTS = {"Ti", "Zr", "Nb", "Hf", "Sc", "Y", "Ta"}
B_SITE_ELEMENTS = {"Mn", "Cr", "Fe", "Co", "Ni", "Cu", "V", "Mo", "W", "Al"}

CORE_COLUMNS = [
    "cif파일 제목",
    "화학식",
    "결정구조",
    "단상 점수",
    "다상 점수",
    "보정 단상 점수",
    "예측결과(신조성 룰)",
    "예측결과(raw hull)",
    "raw_Ehull_위험도",
    "경쟁 유형",
    "예상되는 멀티 phase 조성들",
    "예상 phase 결정구조들",
]

EVIDENCE_COLUMNS = [
    "raw_energy_above_hull_eV_per_atom",
    "config_entropy_kB_per_atom",
    "entropy_temperature_K",
    "entropy_stabilization_eV_per_atom",
    "entropy_weight",
    "entropy_corrected_e_hull_eV_per_atom",
    "single_threshold_eV_per_atom",
    "metastable_threshold_eV_per_atom",
    "A_site_elements",
    "B_site_elements",
    "B_to_A_ratio",
    "laves_AB2_like",
    "calibration_domain",
    "calibration_basis",
    "literature_match",
    "literature_basis",
]

EXTRA_COLUMNS = [
    "chgnet_energy_eV_per_atom",
    "hull_energy_source",
    "space_group",
    "min_distance_A",
    "num_sites",
    "chemical_system",
    "상태",
    "메시지",
]

ALL_COLUMNS = CORE_COLUMNS + EVIDENCE_COLUMNS + EXTRA_COLUMNS
REQUIRED_COLUMNS = CORE_COLUMNS

EXPLANATION_ROWS = [
    ("cif파일 제목", "평가에 사용한 CIF 파일명입니다."),
    ("화학식", "pymatgen이 CIF에서 읽은 reduced formula입니다."),
    ("결정구조", "CIF 구조에서 SpacegroupAnalyzer(symprec=0.1)로 추정한 공간군입니다."),
    ("단상 점수", "raw CHGNet-only E_hull 기준 점수 exp(-E_hull / single_threshold)입니다."),
    ("다상 점수", "1 - 단상 점수입니다. raw hull에서 분해 구동력이 클수록 커집니다."),
    ("보정 단상 점수", "entropy-corrected E_hull 기준 점수입니다."),
    ("예측결과(신조성 룰)", "문헌 직접 매칭 여부를 최종 판정에 쓰지 않고, entropy-corrected E_hull 기준으로 신조성 후보를 분류한 라벨입니다."),
    ("예측결과(raw hull)", "기존 방식과 같은 0 K ordered CHGNet hull 기준 판정입니다."),
    ("raw_Ehull_위험도", "entropy 보정 전 raw E_hull의 위험 구간입니다. raw 값이 높으면 합성 조건 민감성이 커질 수 있습니다."),
    ("경쟁 유형", "분해상이 같은 조성의 polymorph 경쟁인지, 다른 조성 조합으로의 다상 분해인지 구분합니다."),
    ("예상되는 멀티 phase 조성들", "CHGNet-only phase diagram에서 계산된 decomposition product 조합입니다."),
    ("예상 phase 결정구조들", "decomposition product의 공간군과 MP entry id입니다."),
    ("raw_energy_above_hull_eV_per_atom", "CHGNet으로 후보와 MP competing phase를 모두 재계산해 만든 0 K hull distance입니다."),
    ("config_entropy_kB_per_atom", "A/B sublattice 혼합을 가정한 configurational entropy입니다. 단위는 kB/atom입니다."),
    ("entropy_temperature_K", "T*S_config 보정에 사용한 온도입니다. 기본값은 1000 K입니다."),
    ("entropy_stabilization_eV_per_atom", "T * S_config 값입니다. E_hull에서 차감되는 안정화 에너지입니다."),
    ("entropy_weight", "entropy 안정화 항을 어느 정도 반영했는지 나타내는 가중치입니다."),
    ("entropy_corrected_e_hull_eV_per_atom", "raw E_hull - entropy_weight*T*S_config 입니다. 0보다 작으면 0으로 표시합니다."),
    ("single_threshold_eV_per_atom", "단상 후보로 분류하는 E_hull 기준입니다."),
    ("metastable_threshold_eV_per_atom", "고온/엔트로피 안정화 가능 후보로 남기는 상한 기준입니다."),
    ("A_site_elements", "Laves AB2의 A site로 간주한 원소군입니다."),
    ("B_site_elements", "Laves AB2의 B site로 간주한 원소군입니다."),
    ("B_to_A_ratio", "B site 원자수 / A site 원자수입니다. AB2 Laves는 대략 2입니다."),
    ("laves_AB2_like", "A/B site 원소군과 B/A 비율로 판단한 C14 AB2 Laves 후보 여부입니다."),
    ("calibration_domain", "제공 논문으로 보정한 설계공간에 들어가는지 나타냅니다. 신조성에는 문헌 직접 매칭 대신 이 값을 참고합니다."),
    ("calibration_basis", "calibration domain 판단에 사용한 A/B site, B/A ratio, entropy 조건 요약입니다."),
    ("literature_match", "제공된 4개 논문 조성군과 직접 매칭되는지 표시합니다. 신조성 최종 판정에는 사용하지 않습니다."),
    ("literature_basis", "논문에서 확인한 실험적 단상 근거를 요약합니다. 신조성에는 참고용입니다."),
    ("chgnet_energy_eV_per_atom", "CHGNet이 예측한 후보 구조의 원자당 에너지입니다."),
    ("hull_energy_source", "hull 계산에 사용한 에너지 출처입니다."),
    ("space_group", "결정구조와 같은 공간군 값을 영문/진단 컬럼으로 반복 기록합니다."),
    ("min_distance_A", "CIF 내 최단 원자간 거리입니다. 비정상적으로 짧으면 구조 입력 문제 가능성이 있습니다."),
    ("num_sites", "CIF 구조의 site 수입니다."),
    ("chemical_system", "평가에 사용한 원소 조합입니다."),
    ("상태", "평가 성공/주의/실패 상태입니다."),
    ("메시지", "실패 사유 또는 구조상 주의 메시지입니다."),
]


def load_dotenv(path):
    if not path.exists():
        return
    with path.open(encoding="utf-8") as file:
        for line in file:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


def format_float(value, digits=6):
    if value is None:
        return ""
    try:
        if math.isnan(value):
            return ""
    except TypeError:
        pass
    return f"{float(value):.{digits}f}"


def min_interatomic_distance(structure):
    if len(structure) < 2:
        return None
    distances = np.array(structure.distance_matrix, dtype=float)
    np.fill_diagonal(distances, np.inf)
    return float(np.min(distances))


def get_space_group(structure):
    try:
        analyzer = SpacegroupAnalyzer(structure, symprec=0.1)
        return analyzer.get_space_group_symbol()
    except Exception:
        return ""


def load_chgnet_model():
    try:
        from chgnet.model import CHGNet
    except ImportError as exc:
        raise RuntimeError(
            "chgnet is not installed. Install requirements first: "
            "python -m pip install -r requirements.txt"
        ) from exc
    return CHGNet.load()


def predict_energy_per_atom(model, structure):
    prediction = model.predict_structure(structure, task="e")
    energy = prediction["e"]
    if hasattr(energy, "item"):
        return float(energy.item())
    return float(energy)


def build_chgnet_phase_entries(model, mp_entries):
    chgnet_entries = []
    for entry in mp_entries:
        structure = getattr(entry, "structure", None)
        if structure is None:
            continue
        try:
            energy_per_atom = predict_energy_per_atom(model, structure)
        except Exception:
            continue
        chgnet_entries.append(
            ComputedStructureEntry(
                structure=structure,
                energy=energy_per_atom * len(structure),
                entry_id=getattr(entry, "entry_id", None),
            )
        )
    return chgnet_entries


def get_mp_entries(elements, api_key):
    if not api_key:
        raise RuntimeError("MP_API_KEY is not set. Add it to .env or environment variables.")
    try:
        from mp_api.client import MPRester
    except ImportError as exc:
        raise RuntimeError(
            "mp-api is not installed. Install requirements first: "
            "python -m pip install -r requirements.txt"
        ) from exc

    with MPRester(api_key) as mpr:
        return mpr.get_entries_in_chemsys(
            elements,
            additional_criteria={"thermo_types": ["GGA_GGA+U"]},
        )


def decomposition_to_text(decomposition):
    if not decomposition:
        return ""
    parts = []
    for entry, amount in decomposition.items():
        parts.append(f"{entry.composition.reduced_formula} ({amount:.3g})")
    return " + ".join(parts)


def decomposition_structures_to_text(decomposition):
    if not decomposition:
        return ""
    parts = []
    for entry, amount in decomposition.items():
        formula = entry.composition.reduced_formula
        structure = getattr(entry, "structure", None)
        space_group = get_space_group(structure) if structure is not None else ""
        entry_id = getattr(entry, "entry_id", "")
        if space_group:
            parts.append(f"{formula} ({amount:.3g}): {space_group} [{entry_id}]")
        else:
            parts.append(f"{formula} ({amount:.3g}): unknown [{entry_id}]")
    return " + ".join(parts)


def classify_competition(candidate_composition, decomposition):
    if not decomposition:
        return "정보 없음"
    candidate_formula = candidate_composition.reduced_formula
    formulas = [entry.composition.reduced_formula for entry in decomposition]
    unique_formulas = sorted(set(formulas))
    if len(unique_formulas) == 1 and unique_formulas[0] == candidate_formula:
        return "동일 조성 구조경쟁"
    return "다상 분해 가능"


def score_from_e_above_hull(e_above_hull, threshold):
    if e_above_hull is None:
        return None, None, "평가불가"
    single_score = math.exp(-max(e_above_hull, 0.0) / threshold)
    single_score = max(0.0, min(1.0, single_score))
    multi_score = 1.0 - single_score
    prediction = "단상 후보" if e_above_hull <= threshold else "멀티 가능성 높음"
    return single_score, multi_score, prediction


def prediction_with_competition(e_above_hull, threshold, competition_type):
    if e_above_hull is None:
        return "평가불가"
    if e_above_hull <= threshold:
        return "단상 후보"
    if competition_type == "동일 조성 구조경쟁":
        return "단상(다른 결정구조 가능)"
    return "멀티 가능성 높음"


def shannon_entropy(fractions):
    entropy = 0.0
    for fraction in fractions:
        if fraction > 0:
            entropy -= fraction * math.log(fraction)
    return entropy


def laves_entropy_features(composition, temperature_k, entropy_weight):
    amounts = {str(el): float(amount) for el, amount in composition.get_el_amt_dict().items()}
    a_amounts = {el: amt for el, amt in amounts.items() if el in A_SITE_ELEMENTS}
    b_amounts = {el: amt for el, amt in amounts.items() if el in B_SITE_ELEMENTS and el not in a_amounts}
    a_total = sum(a_amounts.values())
    b_total = sum(b_amounts.values())
    total = a_total + b_total

    if a_total <= 0 or b_total <= 0 or total <= 0:
        return {
            "A_site_elements": "",
            "B_site_elements": "",
            "B_to_A_ratio": None,
            "laves_AB2_like": False,
            "config_entropy_kB_per_atom": 0.0,
            "entropy_stabilization_eV_per_atom": 0.0,
            "weighted_entropy_stabilization_eV_per_atom": 0.0,
        }

    a_fractions = [amount / a_total for amount in a_amounts.values()]
    b_fractions = [amount / b_total for amount in b_amounts.values()]
    entropy_kb = (a_total * shannon_entropy(a_fractions) + b_total * shannon_entropy(b_fractions)) / total
    entropy_ev = KB_EV_PER_K * temperature_k * entropy_kb
    b_to_a = b_total / a_total
    laves_like = 1.5 <= b_to_a <= 2.5 and len(a_amounts) >= 1 and len(b_amounts) >= 2

    return {
        "A_site_elements": ", ".join(sorted(a_amounts)),
        "B_site_elements": ", ".join(sorted(b_amounts)),
        "B_to_A_ratio": b_to_a,
        "laves_AB2_like": laves_like,
        "config_entropy_kB_per_atom": entropy_kb if laves_like else 0.0,
        "entropy_stabilization_eV_per_atom": entropy_ev if laves_like else 0.0,
        "weighted_entropy_stabilization_eV_per_atom": entropy_weight * entropy_ev if laves_like else 0.0,
    }


def literature_evidence(cif_name):
    name = cif_name.replace(" ", "")
    if "Ti1-xZrxMn1.0Cr0.85Fe0.1" in name:
        return (
            "1번논문",
            "Ti1-xZrxMn1.0Cr0.85Fe0.1 (x=0.1-0.5): XRD에서 모든 합금이 secondary peak 없는 single C14 Laves phase로 보고됨.",
        )
    if "Ti0.84Zr0.16Mn0.9+xCr0.7Fe0.1" in name or "Ti0.84Zr0.16Mn0.9Cr0.7+yFe0.1" in name:
        return (
            "2번논문",
            "Ti0.84Zr0.16Mn0.9+xCr0.7Fe0.1 및 Ti0.84Zr0.16Mn0.9Cr0.7+yFe0.1: 조성 변화에도 single C14-type Laves phase로 보고됨.",
        )
    if "Ti0.5Zr0.5MnCr" in name or "(Ti0.5Zr0.5)1(Mn0.5Cr0.5)2" in name:
        return (
            "3번논문",
            "(Ti0.5Zr0.5)1(Mn0.5Cr0.5)2: Rietveld/XRD에서 single C14 Laves phase, secondary phase reflection 없음.",
        )
    if "Ti0.5Zr0.5Fe0.66Mn0.66Cr0.66" in name or "(Ti0.5Zr0.5)1(Fe0.33Mn0.33Cr0.33)2" in name:
        return (
            "4번논문",
            "(Ti0.5Zr0.5)1(Fe0.33Mn0.33Cr0.33)2: XRD/Rietveld에서 single C14 Laves phase 및 secondary reflection 없음.",
        )
    return (
        "직접 매칭 없음",
        "제공 논문과 같은 Ti/Zr 기반 C14 AB2 Laves 설계공간인지 entropy 보정값으로 별도 판단.",
    )


def raw_ehull_risk(raw_e_hull):
    if raw_e_hull is None:
        return "평가불가"
    if raw_e_hull <= 0.08:
        return "낮음"
    if raw_e_hull <= 0.15:
        return "중간"
    return "높음"


def calibration_domain_features(entropy, elements):
    allowed_elements = A_SITE_ELEMENTS | B_SITE_ELEMENTS
    in_element_space = set(elements).issubset(allowed_elements)
    entropy_kb = entropy["config_entropy_kB_per_atom"]
    b_to_a = entropy["B_to_A_ratio"]
    laves_like = entropy["laves_AB2_like"]
    in_entropy_range = 0.45 <= entropy_kb <= 1.05
    in_ratio_range = b_to_a is not None and 1.6 <= b_to_a <= 2.4
    in_domain = bool(laves_like and in_element_space and in_entropy_range and in_ratio_range)
    if in_domain:
        domain = "Ti/Zr 기반 C14 AB2 Laves calibration domain"
    elif laves_like:
        domain = "C14 AB2-like, calibration 외삽 영역"
    else:
        domain = "calibration domain 밖"

    basis = (
        f"elements={','.join(elements)}; "
        f"A_site={entropy['A_site_elements']}; B_site={entropy['B_site_elements']}; "
        f"B/A={format_float(b_to_a, 3)}; "
        f"S_config={format_float(entropy_kb, 3)} kB/atom"
    )
    return domain, basis


def screening_prediction(raw_e_hull, corrected_e_hull, threshold, metastable_threshold, laves_like):
    if raw_e_hull is None:
        return "평가불가"
    if raw_e_hull <= threshold:
        return "단상 후보"
    if laves_like and corrected_e_hull <= threshold:
        return "엔트로피 보정 단상 후보"
    if laves_like and corrected_e_hull <= metastable_threshold:
        return "고온 단상 가능 후보"
    if laves_like and corrected_e_hull <= 0.12:
        return "준안정 Laves 후보(계산상 위험 높음)"
    return "멀티 위험 후보"


def evaluate_cif(
    cif_path,
    chgnet_model,
    mp_cache,
    chgnet_hull_cache,
    api_key,
    threshold,
    metastable_threshold,
    entropy_temperature_k,
    entropy_weight,
):
    row = {column: "" for column in ALL_COLUMNS}
    row["cif파일 제목"] = cif_path.name
    row["single_threshold_eV_per_atom"] = format_float(threshold)
    row["metastable_threshold_eV_per_atom"] = format_float(metastable_threshold)
    row["entropy_temperature_K"] = format_float(entropy_temperature_k, 1)
    row["entropy_weight"] = format_float(entropy_weight, 3)

    try:
        structure = Structure.from_file(cif_path)
        if not structure.is_ordered:
            raise RuntimeError("Disordered CIF is not supported yet. Provide an ordered CIF.")

        composition = structure.composition
        formula = composition.reduced_formula
        elements = sorted([el.symbol for el in composition.elements])
        chemical_system = "-".join(elements)
        min_dist = min_interatomic_distance(structure)

        row["화학식"] = formula
        row["결정구조"] = get_space_group(structure)
        row["space_group"] = row["결정구조"]
        row["min_distance_A"] = format_float(min_dist, 4)
        row["num_sites"] = str(len(structure))
        row["chemical_system"] = chemical_system

        if min_dist is not None and min_dist < 0.7:
            row["상태"] = "주의"
            row["메시지"] = f"Very short interatomic distance: {min_dist:.3f} A"

        energy_per_atom = predict_energy_per_atom(chgnet_model, structure)
        row["chgnet_energy_eV_per_atom"] = format_float(energy_per_atom)

        if chemical_system not in mp_cache:
            mp_cache[chemical_system] = get_mp_entries(elements, api_key)
        mp_entries = mp_cache[chemical_system]
        if not mp_entries:
            raise RuntimeError(f"No Materials Project entries found for {chemical_system}.")

        if chemical_system not in chgnet_hull_cache:
            chgnet_hull_cache[chemical_system] = build_chgnet_phase_entries(chgnet_model, mp_entries)
        chgnet_entries = chgnet_hull_cache[chemical_system]
        if not chgnet_entries:
            raise RuntimeError(f"No MP structures available for CHGNet hull: {chemical_system}.")

        candidate_entry = ComputedStructureEntry(
            structure=structure,
            energy=energy_per_atom * len(structure),
            entry_id=cif_path.stem,
        )
        phase_diagram = PhaseDiagram(chgnet_entries + [candidate_entry])
        decomposition, raw_e_hull = phase_diagram.get_decomp_and_e_above_hull(candidate_entry)

        entropy = laves_entropy_features(composition, entropy_temperature_k, entropy_weight)
        corrected_e_hull = max(0.0, raw_e_hull - entropy["weighted_entropy_stabilization_eV_per_atom"])
        literature_match, literature_basis = literature_evidence(cif_path.name)
        calibration_domain, calibration_basis = calibration_domain_features(entropy, elements)

        competition_type = classify_competition(composition, decomposition)
        single_score, multi_score, _ = score_from_e_above_hull(raw_e_hull, threshold)
        corrected_single_score, _, _ = score_from_e_above_hull(corrected_e_hull, threshold)

        row["단상 점수"] = format_float(single_score)
        row["다상 점수"] = format_float(multi_score)
        row["보정 단상 점수"] = format_float(corrected_single_score)
        row["예측결과(raw hull)"] = prediction_with_competition(raw_e_hull, threshold, competition_type)
        row["예측결과(신조성 룰)"] = screening_prediction(
            raw_e_hull,
            corrected_e_hull,
            threshold,
            metastable_threshold,
            entropy["laves_AB2_like"],
        )
        row["raw_Ehull_위험도"] = raw_ehull_risk(raw_e_hull)
        row["경쟁 유형"] = competition_type
        row["예상되는 멀티 phase 조성들"] = decomposition_to_text(decomposition)
        row["예상 phase 결정구조들"] = decomposition_structures_to_text(decomposition)
        row["raw_energy_above_hull_eV_per_atom"] = format_float(raw_e_hull)
        row["config_entropy_kB_per_atom"] = format_float(entropy["config_entropy_kB_per_atom"])
        row["entropy_stabilization_eV_per_atom"] = format_float(entropy["entropy_stabilization_eV_per_atom"])
        row["entropy_corrected_e_hull_eV_per_atom"] = format_float(corrected_e_hull)
        row["A_site_elements"] = entropy["A_site_elements"]
        row["B_site_elements"] = entropy["B_site_elements"]
        row["B_to_A_ratio"] = format_float(entropy["B_to_A_ratio"], 4)
        row["laves_AB2_like"] = "TRUE" if entropy["laves_AB2_like"] else "FALSE"
        row["calibration_domain"] = calibration_domain
        row["calibration_basis"] = calibration_basis
        row["literature_match"] = literature_match
        row["literature_basis"] = literature_basis
        row["hull_energy_source"] = "CHGNet-only hull from MP structures"

        if not row["상태"]:
            row["상태"] = "성공"

    except Exception as exc:
        if not row["상태"]:
            row["상태"] = "실패"
        row["메시지"] = str(exc)
        if not row["예측결과(신조성 룰)"]:
            row["예측결과(신조성 룰)"] = "평가불가"
        if not row["예측결과(raw hull)"]:
            row["예측결과(raw hull)"] = "평가불가"

    return row


def collect_cif_files(input_dir, max_files=None):
    if input_dir.is_file():
        cif_files = [input_dir] if input_dir.suffix.lower() == ".cif" else []
    else:
        cif_files = sorted(input_dir.glob("*.cif"))
    if max_files is not None:
        cif_files = cif_files[:max_files]
    return cif_files


def write_csv(rows, result_dir):
    result_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = result_dir / f"phase_result_{timestamp}.csv"
    with output_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=ALL_COLUMNS)
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, "") for column in ALL_COLUMNS})
    return output_path


def style_sheet(ws):
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:100])
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 55)


def write_excel(rows, result_dir):
    result_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = result_dir / f"phase_result_{timestamp}.xlsx"

    wb = Workbook()
    ws_info = wb.active
    ws_info.title = "항목 설명"
    ws_info.append(["항목", "설명"])
    for item, description in EXPLANATION_ROWS:
        ws_info.append([item, description])
    style_sheet(ws_info)

    ws_result = wb.create_sheet("phase_result")
    ws_result.append(ALL_COLUMNS)
    for row in rows:
        ws_result.append([row.get(column, "") for column in ALL_COLUMNS])
    style_sheet(ws_result)

    ws_lit = wb.create_sheet("문헌 근거")
    ws_lit.append(["논문", "대상 조성", "실험 단상 근거 요약"])
    ws_lit.append(["1번논문", "Ti1-xZrxMn1.0Cr0.85Fe0.1 (x=0.1-0.5)", "XRD에서 모든 합금이 secondary peak 없는 single C14 Laves phase로 보고됨."])
    ws_lit.append(["2번논문", "Ti0.84Zr0.16Mn0.9+xCr0.7Fe0.1, Ti0.84Zr0.16Mn0.9Cr0.7+yFe0.1", "조성 변화에도 single C14-type Laves phase로 결정화, cycling 후에도 대표 조성의 C14 구조 유지 보고."])
    ws_lit.append(["3번논문", "(Ti0.5Zr0.5)1(Mn0.5Cr0.5)2", "Rietveld/XRD에서 single C14 Laves phase로 분석되며 secondary phase reflection 없음."])
    ws_lit.append(["4번논문", "(Ti0.5Zr0.5)1(Fe0.33Mn0.33Cr0.33)2", "XRD/Rietveld에서 single C14 Laves phase 및 secondary reflection 없음."])
    style_sheet(ws_lit)

    wb.save(output_path)
    return output_path


def parse_args():
    parser = argparse.ArgumentParser(
        description="Evaluate substituted CIF files for single/multi phase tendency."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT_DIR, help="CIF file or folder.")
    parser.add_argument("--result-dir", type=Path, default=DEFAULT_RESULT_DIR)
    parser.add_argument("--max-files", type=int, default=None)
    parser.add_argument(
        "--single-threshold",
        type=float,
        default=0.03,
        help="Energy-above-hull threshold in eV/atom for single-phase candidate.",
    )
    parser.add_argument(
        "--metastable-threshold",
        type=float,
        default=0.08,
        help="Corrected E_hull upper bound for high-temperature single-phase candidates.",
    )
    parser.add_argument(
        "--entropy-temperature",
        type=float,
        default=DEFAULT_ENTROPY_TEMPERATURE_K,
        help="Temperature in K for T*S_config stabilization.",
    )
    parser.add_argument(
        "--entropy-weight",
        type=float,
        default=DEFAULT_ENTROPY_WEIGHT,
        help="Fraction of configurational entropy stabilization applied to corrected E_hull.",
    )
    return parser.parse_args()


def main():
    load_dotenv(ROOT / ".env")
    args = parse_args()

    input_dir = args.input if args.input.is_absolute() else ROOT / args.input
    result_dir = args.result_dir if args.result_dir.is_absolute() else ROOT / args.result_dir
    cif_files = collect_cif_files(input_dir, args.max_files)

    if not cif_files:
        print(f"No CIF files found in {input_dir}")
        print("No result file was written.")
        return

    chgnet_model = load_chgnet_model()
    api_key = os.getenv("MP_API_KEY")
    mp_cache = {}
    chgnet_hull_cache = {}

    rows = []
    for idx, cif_path in enumerate(cif_files, start=1):
        print(f"[{idx}/{len(cif_files)}] Evaluating {cif_path.name}")
        rows.append(
            evaluate_cif(
                cif_path,
                chgnet_model,
                mp_cache,
                chgnet_hull_cache,
                api_key,
                args.single_threshold,
                args.metastable_threshold,
                args.entropy_temperature,
                args.entropy_weight,
            )
        )

    csv_path = write_csv(rows, result_dir)
    xlsx_path = write_excel(rows, result_dir)
    print(f"Result CSV written to {csv_path}")
    print(f"Result Excel written to {xlsx_path}")


if __name__ == "__main__":
    main()
